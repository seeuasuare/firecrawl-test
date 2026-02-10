#!/usr/bin/env python3
"""
Site audit script — checks each company website for:
  1. Accessibility (opens or not, status code, redirect)
  2. Thematic relevance ("children's psychological center")
  3. Recent activity (last 3 months)

Usage:
    pip install openpyxl aiohttp beautifulsoup4
    python audit_sites.py

Input:  companies_final.xlsx
Output: companies_audit.xlsx  (with new columns: Status, Theme Match, Activity)
"""

import asyncio
import re
import time
import warnings
from datetime import datetime, timedelta
from urllib.parse import urljoin, urlparse

import aiohttp
import openpyxl
from openpyxl.styles import Font, PatternFill
from bs4 import BeautifulSoup, XMLParsedAsHTMLWarning

warnings.filterwarnings("ignore", category=XMLParsedAsHTMLWarning)
warnings.filterwarnings("ignore", category=DeprecationWarning)

# --------------- Configuration ---------------

CONCURRENCY = 25
TIMEOUT = 12
INPUT_FILE = 'companies_final.xlsx'
OUTPUT_FILE = 'companies_audit.xlsx'

# How far back to consider "recent" activity
MONTHS_BACK = 3

# --------------- Theme keywords ---------------

# Primary keywords (strong signal — directly about child psychology)
PRIMARY_KEYWORDS = [
    'детский психолог', 'детская психолог', 'детских психолог',
    'нейропсихолог', 'нейро-психолог',
    'детский центр', 'детский развива',
    'логопед', 'дефектолог',
    'детский нейро', 'нейрокоррекц',
    'психолог для детей', 'психолог для ребенк', 'психолог для ребёнк',
    'психологическ центр', 'психологический центр',
    'коррекционн', 'коррекция',
    'развити ребенк', 'развити ребёнк', 'развитие детей',
    'aba-терап', 'аба-терап', 'аба терап',
    'сенсорная интеграц', 'канистерап',
    'детский невролог', 'детский психиатр',
    'песочная терап', 'арт-терап', 'игровая терап',
    'задержк развит', 'зпр', 'зрр', 'сдвг', 'аутизм', 'рас ',
    'дислекси', 'дисграф', 'алали', 'заикани',
]

# Secondary keywords (weaker signal — could be any psychology)
SECONDARY_KEYWORDS = [
    'психолог', 'психотерап', 'консультац',
    'тревожност', 'эмоциональн',
]

# Anti-keywords (if these dominate, it's probably NOT a child psych center)
ANTI_KEYWORDS = [
    'автосервис', 'автомобил', 'ремонт квартир', 'натяжн потолк',
    'юридическ', 'бухгалтер', 'строительств', 'недвижимост',
    'стоматолог', 'кулинар', 'доставка еды',
]

# --------------- Date detection ---------------

# Russian month names for date parsing
RUSSIAN_MONTHS = {
    'январ': 1, 'феврал': 2, 'март': 3, 'апрел': 4,
    'мая': 5, 'мае': 5, 'май': 5, 'июн': 6, 'июл': 7,
    'август': 8, 'сентябр': 9, 'октябр': 10,
    'ноябр': 11, 'декабр': 12,
}

# Regex for dates like "15 января 2026", "15.01.2026", "2026-01-15"
DATE_PATTERNS = [
    # 15 января 2026
    re.compile(r'(\d{1,2})\s+(' + '|'.join(RUSSIAN_MONTHS.keys()) + r')[а-яё]*\s+(\d{4})', re.IGNORECASE),
    # 15.01.2026 or 15/01/2026
    re.compile(r'(\d{1,2})[./](\d{1,2})[./](\d{4})'),
    # 2026-01-15
    re.compile(r'(\d{4})-(\d{2})-(\d{2})'),
]

# Copyright year pattern
COPYRIGHT_RE = re.compile(r'©\s*(?:\d{4}\s*[-–—]\s*)?(\d{4})', re.IGNORECASE)
COPYRIGHT_RE2 = re.compile(r'copyright\s*(?:©)?\s*(?:\d{4}\s*[-–—]\s*)?(\d{4})', re.IGNORECASE)


def parse_russian_date(match, pattern_idx):
    """Convert regex match to datetime."""
    try:
        if pattern_idx == 0:  # "15 января 2026"
            day = int(match.group(1))
            month_text = match.group(2).lower()
            year = int(match.group(3))
            month = None
            for key, val in RUSSIAN_MONTHS.items():
                if month_text.startswith(key):
                    month = val
                    break
            if month and 2000 <= year <= 2030:
                return datetime(year, month, day)
        elif pattern_idx == 1:  # "15.01.2026"
            day, month, year = int(match.group(1)), int(match.group(2)), int(match.group(3))
            if 2000 <= year <= 2030 and 1 <= month <= 12:
                return datetime(year, month, day)
        elif pattern_idx == 2:  # "2026-01-15"
            year, month, day = int(match.group(1)), int(match.group(2)), int(match.group(3))
            if 2000 <= year <= 2030 and 1 <= month <= 12:
                return datetime(year, month, day)
    except (ValueError, IndexError):
        pass
    return None


def find_dates_in_text(text):
    """Find all recognizable dates in text."""
    dates = []
    for idx, pattern in enumerate(DATE_PATTERNS):
        for match in pattern.finditer(text):
            dt = parse_russian_date(match, idx)
            if dt:
                dates.append(dt)
    return dates


def find_copyright_year(text):
    """Extract copyright year from page."""
    for pattern in [COPYRIGHT_RE, COPYRIGHT_RE2]:
        match = pattern.search(text)
        if match:
            year = int(match.group(1))
            if 2000 <= year <= 2030:
                return year
    return None


# --------------- Audit logic ---------------

async def fetch_with_info(session, url):
    """Fetch URL and return (status_code, final_url, html, headers)."""
    try:
        async with session.get(
            url,
            timeout=aiohttp.ClientTimeout(total=TIMEOUT),
            ssl=False,
            allow_redirects=True,
        ) as resp:
            status = resp.status
            final_url = str(resp.url)
            headers = dict(resp.headers)
            html = ''
            if status == 200:
                ct = resp.headers.get('Content-Type', '')
                if 'text' in ct or 'html' in ct or not ct:
                    html = await resp.text(errors='replace')
            return status, final_url, html, headers
    except aiohttp.ClientConnectorError:
        return 0, url, '', {}
    except asyncio.TimeoutError:
        return -1, url, '', {}
    except Exception as e:
        return -2, url, '', {}


def check_theme(html):
    """Check if site is about children's psychology. Returns (is_match, confidence, details)."""
    if not html:
        return None, 0, 'Нет данных'

    text = BeautifulSoup(html, 'html.parser').get_text(' ', strip=True).lower()

    # Check anti-keywords first
    anti_count = sum(1 for kw in ANTI_KEYWORDS if kw in text)
    if anti_count >= 2:
        return False, 0, 'Другая тематика'

    primary_found = [kw for kw in PRIMARY_KEYWORDS if kw in text]
    secondary_found = [kw for kw in SECONDARY_KEYWORDS if kw in text]

    if len(primary_found) >= 3:
        return True, 100, f'Точное совпадение ({len(primary_found)} ключевых слов)'
    elif len(primary_found) >= 1:
        confidence = min(50 + len(primary_found) * 20 + len(secondary_found) * 10, 100)
        return True, confidence, f'Совпадение ({len(primary_found)} осн. + {len(secondary_found)} доп.)'
    elif len(secondary_found) >= 2:
        return True, 40, f'Возможное совпадение ({len(secondary_found)} доп. слов)'
    elif len(secondary_found) == 1:
        return None, 20, 'Слабое совпадение (1 общее слово)'
    else:
        return False, 0, 'Не соответствует тематике'


def check_activity(html, headers):
    """Check for recent site activity. Returns (is_active, details)."""
    now = datetime.now()
    cutoff = now - timedelta(days=MONTHS_BACK * 30)

    results = []

    # 1) Check Last-Modified header
    last_mod = headers.get('Last-Modified', '')
    if last_mod:
        try:
            from email.utils import parsedate_to_datetime
            lm_date = parsedate_to_datetime(last_mod)
            lm_date = lm_date.replace(tzinfo=None)
            if lm_date >= cutoff:
                results.append(f'Last-Modified: {lm_date.strftime("%d.%m.%Y")}')
        except Exception:
            pass

    if not html:
        if results:
            return True, '; '.join(results)
        return None, 'Нет данных'

    text = BeautifulSoup(html, 'html.parser').get_text(' ', strip=True)

    # 2) Find dates in content
    dates = find_dates_in_text(text)
    recent_dates = [d for d in dates if d >= cutoff and d <= now + timedelta(days=30)]
    if recent_dates:
        latest = max(recent_dates)
        results.append(f'Дата на сайте: {latest.strftime("%d.%m.%Y")} ({len(recent_dates)} свежих дат)')

    # 3) Check copyright year
    cr_year = find_copyright_year(text)
    current_year = now.year
    if cr_year:
        if cr_year >= current_year:
            results.append(f'Copyright {cr_year}')
        elif cr_year == current_year - 1:
            results.append(f'Copyright {cr_year} (прошлый год)')

    if results:
        return True, '; '.join(results)

    # 4) If we found old dates but nothing recent
    if dates:
        oldest = min(dates)
        newest = max(dates)
        return False, f'Последняя дата: {newest.strftime("%d.%m.%Y")}'

    if cr_year and cr_year < current_year - 1:
        return False, f'Copyright {cr_year} (устаревший)'

    return None, 'Нет признаков активности'


async def audit_one(session, url):
    """Audit a single website."""
    if not url or url == 'None':
        return {
            'status': 'Нет URL',
            'status_code': 0,
            'theme': None,
            'theme_confidence': 0,
            'theme_detail': 'Нет URL',
            'active': None,
            'active_detail': 'Нет URL',
        }

    if not url.startswith('http'):
        url = 'https://' + url

    status_code, final_url, html, headers = await fetch_with_info(session, url)

    # 1) Accessibility
    if status_code == 200:
        status = 'OK'
    elif status_code == 0:
        status = 'Не открывается (ошибка соединения)'
    elif status_code == -1:
        status = 'Не открывается (таймаут)'
    elif status_code == -2:
        status = 'Не открывается (ошибка)'
    elif status_code == 403:
        status = 'Заблокирован (403)'
    elif status_code == 404:
        status = 'Не найден (404)'
    elif status_code == 500:
        status = 'Ошибка сервера (500)'
    elif status_code == 502:
        status = 'Ошибка сервера (502)'
    elif status_code == 503:
        status = 'Недоступен (503)'
    elif 300 <= status_code < 400:
        status = f'Редирект ({status_code})'
    else:
        status = f'Ошибка ({status_code})'

    # 2) Theme check
    theme, theme_conf, theme_detail = check_theme(html)

    # 3) Activity check
    active, active_detail = check_activity(html, headers)

    return {
        'status': status,
        'status_code': status_code,
        'theme': theme,
        'theme_confidence': theme_conf,
        'theme_detail': theme_detail,
        'active': active,
        'active_detail': active_detail,
    }


# --------------- Main ---------------

async def main():
    print(f"Loading {INPUT_FILE}...")
    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb.active
    total = ws.max_row - 1

    # Add new header columns
    ws.cell(1, 6).value = 'Сайт работает?'
    ws.cell(1, 7).value = 'Тематика'
    ws.cell(1, 8).value = 'Уверенность %'
    ws.cell(1, 9).value = 'Детали тематики'
    ws.cell(1, 10).value = 'Активность'
    ws.cell(1, 11).value = 'Детали активности'

    # Bold headers
    for col in range(6, 12):
        ws.cell(1, col).font = Font(bold=True)

    rows = []
    for row_idx in range(2, ws.max_row + 1):
        url = ws.cell(row_idx, 3).value or ws.cell(row_idx, 2).value or ''
        rows.append((row_idx, str(url)))

    print(f"Auditing {total} sites...\n")

    stats = {'ok': 0, 'down': 0, 'theme_yes': 0, 'theme_no': 0,
             'active': 0, 'inactive': 0}
    processed = 0

    connector = aiohttp.TCPConnector(limit=CONCURRENCY, ssl=False)
    headers = {
        'User-Agent': (
            'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
            'AppleWebKit/537.36 (KHTML, like Gecko) '
            'Chrome/120.0.0.0 Safari/537.36'
        ),
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
        'Accept-Language': 'ru-RU,ru;q=0.9,en;q=0.8',
    }

    # Color fills
    green = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    red = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    yellow = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')

    async with aiohttp.ClientSession(connector=connector, headers=headers) as session:
        sem = asyncio.Semaphore(CONCURRENCY)

        async def process(row_idx, url):
            nonlocal processed
            async with sem:
                try:
                    async with asyncio.timeout(30):
                        result = await audit_one(session, url)
                except (asyncio.TimeoutError, TimeoutError):
                    result = {
                        'status': 'Не открывается (таймаут)',
                        'status_code': 0,
                        'theme': None, 'theme_confidence': 0,
                        'theme_detail': 'Нет данных',
                        'active': None, 'active_detail': 'Нет данных',
                    }

                # Write results
                ws.cell(row_idx, 6).value = result['status']
                if result['status_code'] == 200:
                    ws.cell(row_idx, 6).fill = green
                    stats['ok'] += 1
                elif result['status_code'] > 0:
                    ws.cell(row_idx, 6).fill = yellow
                else:
                    ws.cell(row_idx, 6).fill = red
                    stats['down'] += 1

                # Theme
                if result['theme'] is True:
                    ws.cell(row_idx, 7).value = 'Да'
                    ws.cell(row_idx, 7).fill = green
                    stats['theme_yes'] += 1
                elif result['theme'] is False:
                    ws.cell(row_idx, 7).value = 'Нет'
                    ws.cell(row_idx, 7).fill = red
                    stats['theme_no'] += 1
                else:
                    ws.cell(row_idx, 7).value = 'Неясно'
                    ws.cell(row_idx, 7).fill = yellow

                ws.cell(row_idx, 8).value = result['theme_confidence']
                ws.cell(row_idx, 9).value = result['theme_detail']

                # Activity
                if result['active'] is True:
                    ws.cell(row_idx, 10).value = 'Активен'
                    ws.cell(row_idx, 10).fill = green
                    stats['active'] += 1
                elif result['active'] is False:
                    ws.cell(row_idx, 10).value = 'Неактивен'
                    ws.cell(row_idx, 10).fill = red
                    stats['inactive'] += 1
                else:
                    ws.cell(row_idx, 10).value = 'Неизвестно'
                    ws.cell(row_idx, 10).fill = yellow

                ws.cell(row_idx, 11).value = result['active_detail']

                processed += 1
                if processed % 10 == 0 or processed == total:
                    pct = processed * 100 // total
                    bar = '\u2588' * (pct // 5) + '\u2591' * (20 - pct // 5)
                    print(f"\r  [{bar}] {processed}/{total}", end='', flush=True)

        tasks = [process(r, u) for r, u in rows]
        await asyncio.gather(*tasks)

    # Auto-adjust column widths
    for col in range(6, 12):
        max_len = len(str(ws.cell(1, col).value))
        for row in range(2, min(50, ws.max_row + 1)):
            val = ws.cell(row, col).value
            if val:
                max_len = max(max_len, min(len(str(val)), 50))
        ws.column_dimensions[chr(64 + col)].width = max_len + 2

    print(f"\n\n{'='*55}")
    print(f"  AUDIT COMPLETE — {total} sites")
    print(f"{'='*55}")
    print(f"  Доступность:")
    print(f"    Работает:         {stats['ok']}")
    print(f"    Не работает:      {stats['down']}")
    print(f"    Другое:           {total - stats['ok'] - stats['down']}")
    print(f"")
    print(f"  Тематика 'детский психологический центр':")
    print(f"    Соответствует:    {stats['theme_yes']}")
    print(f"    Не соответствует: {stats['theme_no']}")
    print(f"    Неясно:           {total - stats['theme_yes'] - stats['theme_no']}")
    print(f"")
    print(f"  Активность (последние {MONTHS_BACK} мес.):")
    print(f"    Активен:          {stats['active']}")
    print(f"    Неактивен:        {stats['inactive']}")
    print(f"    Неизвестно:       {total - stats['active'] - stats['inactive']}")
    print(f"{'='*55}")

    wb.save(OUTPUT_FILE)
    print(f"\nResults saved to: {OUTPUT_FILE}")


if __name__ == '__main__':
    start = time.time()
    asyncio.run(main())
    elapsed = time.time() - start
    m, s = divmod(int(elapsed), 60)
    print(f"Total time: {m}m {s}s")
