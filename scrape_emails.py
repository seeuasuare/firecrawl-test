#!/usr/bin/env python3
"""
Scrape email addresses from a list of company websites.

Usage:
    pip install openpyxl aiohttp beautifulsoup4
    python scrape_emails.py

Input:  companies.xlsx  (with columns: Name, Domain, URL, Description, Emails)
Output: companies_with_emails.xlsx  (same file with Emails column filled in)
"""

import asyncio
import re
import time
from urllib.parse import urljoin, urlparse

import aiohttp
import openpyxl
from bs4 import BeautifulSoup

# --------------- Configuration ---------------

CONCURRENCY = 30        # parallel requests
TIMEOUT = 15            # seconds per request
INPUT_FILE = 'companies.xlsx'
OUTPUT_FILE = 'companies_with_emails.xlsx'

# --------------- Email extraction ---------------

EMAIL_RE = re.compile(
    r'[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}',
    re.IGNORECASE,
)

CONTACT_PATHS = [
    '/contacts', '/contact', '/kontakty', '/about',
    '/about-us', '/o-nas', '/o-kompanii', '/kontakt',
    '/contact-us', '/svyaz', '/feedback',
]

JUNK_DOMAINS = {
    'example.com', 'email.com', 'domain.com', 'test.com',
    'sentry.io', 'sentry-next.wixpress.com', 'wixpress.com',
    'yoursite.com', 'yourdomain.com', 'site.com',
    'w3.org', 'schema.org', 'gravatar.com', 'wordpress.org',
    'googleusercontent.com',
}

JUNK_PREFIXES = (
    'noreply', 'no-reply', 'support@wix', 'support@tilda',
    'webpack', 'grunt', 'gulp', 'postmaster', 'mailer-daemon',
)

JUNK_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.gif', '.svg', '.webp', '.css', '.js'}


def is_valid_email(email: str) -> bool:
    """Filter out junk/invalid emails."""
    email = email.lower().strip()
    if len(email) > 100 or len(email) < 5:
        return False
    if '@' not in email:
        return False
    local, domain = email.rsplit('@', 1)
    if domain in JUNK_DOMAINS:
        return False
    if email.startswith(JUNK_PREFIXES):
        return False
    for ext in JUNK_EXTENSIONS:
        if email.endswith(ext) or ext in domain:
            return False
    if re.search(r'\d{5,}', email):
        return False
    if '..' in email:
        return False
    if not re.match(r'^[a-z0-9._%+\-]+@[a-z0-9.\-]+\.[a-z]{2,}$', email):
        return False
    return True


def extract_emails(html: str) -> set:
    """Extract email addresses from HTML content."""
    emails = set()

    # 1) mailto: links (most reliable)
    soup = BeautifulSoup(html, 'html.parser')
    for a_tag in soup.find_all('a', href=True):
        href = a_tag['href']
        if 'mailto:' in href:
            raw = href.split('mailto:')[1].split('?')[0].split('#')[0].strip()
            if is_valid_email(raw):
                emails.add(raw.lower())

    # 2) Regex across the full HTML
    for match in EMAIL_RE.findall(html):
        if is_valid_email(match):
            emails.add(match.lower())

    return emails


# --------------- Fetching ---------------

async def fetch_page(session: aiohttp.ClientSession, url: str) -> str:
    """Fetch a single URL, return HTML or empty string on error."""
    try:
        async with session.get(
            url,
            timeout=aiohttp.ClientTimeout(total=TIMEOUT),
            ssl=False,
            allow_redirects=True,
        ) as resp:
            if resp.status == 200:
                ct = resp.headers.get('Content-Type', '')
                if 'text' in ct or 'html' in ct or not ct:
                    return await resp.text(errors='replace')
    except Exception:
        pass
    return ''


async def scrape_one(session: aiohttp.ClientSession, url: str, name: str) -> set:
    """Scrape emails from one website: main page + contact pages."""
    emails = set()

    if not url or url == 'None':
        return emails
    if not url.startswith('http'):
        url = 'https://' + url

    parsed = urlparse(url)
    base = f"{parsed.scheme}://{parsed.netloc}"

    # 1) Main page
    html = await fetch_page(session, url)
    if html:
        emails.update(extract_emails(html))

        # 2) Find contact-like links on the page
        soup = BeautifulSoup(html, 'html.parser')
        checked = set()
        for a_tag in soup.find_all('a', href=True):
            href_lower = a_tag['href'].lower()
            text_lower = (a_tag.get_text() or '').lower()
            keywords = ['contact', 'kontakt', 'контакт', 'связ', 'о нас',
                        'about', 'обратн', 'напис']
            if any(kw in href_lower or kw in text_lower for kw in keywords):
                full_url = urljoin(base, a_tag['href'])
                if urlparse(full_url).netloc == parsed.netloc and full_url not in checked:
                    checked.add(full_url)
                    contact_html = await fetch_page(session, full_url)
                    if contact_html:
                        emails.update(extract_emails(contact_html))

    # 3) Try common contact page paths
    for path in CONTACT_PATHS:
        contact_url = base + path
        contact_html = await fetch_page(session, contact_url)
        if contact_html:
            emails.update(extract_emails(contact_html))

    return emails


# --------------- Main ---------------

async def main():
    print(f"Loading {INPUT_FILE}...")
    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb.active
    total = ws.max_row - 1
    print(f"Found {total} companies to scrape\n")

    rows = []
    for row_idx in range(2, ws.max_row + 1):
        url = ws.cell(row_idx, 3).value   # Column C = URL
        domain = ws.cell(row_idx, 2).value # Column B = Domain
        name = ws.cell(row_idx, 1).value or ''
        target = url or domain or ''
        rows.append((row_idx, str(target), str(name)))

    found_count = 0
    failed_count = 0
    processed = 0

    connector = aiohttp.TCPConnector(limit=CONCURRENCY, ssl=False)
    headers = {
        'User-Agent': (
            'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
            'AppleWebKit/537.36 (KHTML, like Gecko) '
            'Chrome/120.0.0.0 Safari/537.36'
        ),
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
        'Accept-Language': 'ru-RU,ru;q=0.9,en;q=0.8',
    }

    async with aiohttp.ClientSession(connector=connector, headers=headers) as session:
        sem = asyncio.Semaphore(CONCURRENCY)

        async def process(row_idx: int, url: str, name: str):
            nonlocal found_count, failed_count, processed
            async with sem:
                emails = await scrape_one(session, url, name)
                if emails:
                    ws.cell(row_idx, 5).value = ', '.join(sorted(emails))
                    found_count += 1
                processed += 1

                if processed % 25 == 0 or processed == total:
                    pct = processed * 100 // total
                    bar = '█' * (pct // 5) + '░' * (20 - pct // 5)
                    print(f"\r  [{bar}] {processed}/{total}  "
                          f"emails found: {found_count}", end='', flush=True)

        tasks = [process(r, u, n) for r, u, n in rows]
        await asyncio.gather(*tasks)

    print(f"\n\n{'='*50}")
    print(f"  Done! Processed: {total} companies")
    print(f"  Emails found:    {found_count}")
    print(f"  No emails:       {total - found_count}")
    print(f"{'='*50}")

    wb.save(OUTPUT_FILE)
    print(f"\nResults saved to: {OUTPUT_FILE}")


if __name__ == '__main__':
    start = time.time()
    asyncio.run(main())
    elapsed = time.time() - start
    m, s = divmod(int(elapsed), 60)
    print(f"Total time: {m}m {s}s")
