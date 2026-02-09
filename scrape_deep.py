#!/usr/bin/env python3
"""
Deep email scraper — second pass for sites where no email was found.

Reads companies_with_emails.xlsx, takes only rows WITHOUT emails,
and does a deeper crawl with timeouts per site so it never hangs.

Usage:
    pip install openpyxl aiohttp beautifulsoup4
    python scrape_deep.py

Input:  companies_with_emails.xlsx
Output: companies_final.xlsx
"""

import asyncio
import html as html_module
import re
import time
import warnings
from urllib.parse import urljoin, urlparse

import aiohttp
import openpyxl
from bs4 import BeautifulSoup, XMLParsedAsHTMLWarning

# Suppress noisy warnings
warnings.filterwarnings("ignore", category=XMLParsedAsHTMLWarning)
warnings.filterwarnings("ignore", category=DeprecationWarning)

# --------------- Configuration ---------------

CONCURRENCY = 25
TIMEOUT = 10             # seconds per single page request
SITE_TIMEOUT = 45        # max seconds per entire site
MAX_PAGES_PER_SITE = 15  # max pages to check per site
INPUT_FILE = 'companies_with_emails.xlsx'
OUTPUT_FILE = 'companies_final.xlsx'

# --------------- Email extraction ---------------

EMAIL_RE = re.compile(
    r'[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}',
    re.IGNORECASE,
)

CONTACT_PATHS = [
    '/contacts', '/contact', '/kontakty', '/about',
    '/about-us', '/o-nas', '/o-kompanii', '/kontakt',
    '/contact-us', '/svyaz', '/feedback', '/team',
    '/specialists', '/specialisty', '/nashi-specialisty',
    '/komanda', '/o-centre', '/o-klinike', '/o-tsentre',
    '/privacy', '/privacy-policy', '/policy',
    '/politika-konfidencialnosti', '/rekvizity',
]

JUNK_DOMAINS = {
    'example.com', 'email.com', 'domain.com', 'test.com',
    'sentry.io', 'sentry-next.wixpress.com', 'wixpress.com',
    'yoursite.com', 'yourdomain.com', 'site.com',
    'w3.org', 'schema.org', 'gravatar.com', 'wordpress.org',
    'googleusercontent.com', 'tinymce.com', 'jsdelivr.net',
    'cloudflare.com', 'googleapis.com', 'gstatic.com',
    'wordpress.com', 'wp.com', 'jquery.com',
}

JUNK_PREFIXES = (
    'noreply', 'no-reply', 'support@wix', 'support@tilda',
    'webpack', 'grunt', 'gulp', 'postmaster', 'mailer-daemon',
    'admin@wordpress', 'info@starter',
)

JUNK_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.gif', '.svg', '.webp', '.css', '.js', '.woff'}


def is_valid_email(email: str) -> bool:
    email = email.lower().strip()
    if len(email) > 100 or len(email) < 5:
        return False
    if '@' not in email:
        return False
    local, domain = email.rsplit('@', 1)
    if domain in JUNK_DOMAINS:
        return False
    if email.startswith(JUNK_PREFIXES):
        return False
    for ext in JUNK_EXTENSIONS:
        if email.endswith(ext) or ext in domain:
            return False
    if re.search(r'\d{5,}', email):
        return False
    if '..' in email:
        return False
    if not re.match(r'^[a-z0-9._%+\-]+@[a-z0-9.\-]+\.[a-z]{2,}$', email):
        return False
    return True


def extract_emails(html_text: str) -> set:
    """Extract emails from HTML — deep version."""
    emails = set()

    decoded = html_module.unescape(html_text)

    # 1) mailto: links
    soup = BeautifulSoup(decoded, 'html.parser')
    for a_tag in soup.find_all('a', href=True):
        href = a_tag['href']
        if 'mailto:' in href:
            raw = href.split('mailto:')[1].split('?')[0].split('#')[0].strip()
            if is_valid_email(raw):
                emails.add(raw.lower())

    # 2) Regex on full HTML
    for match in EMAIL_RE.findall(decoded):
        if is_valid_email(match):
            emails.add(match.lower())

    # 3) data-attributes
    for tag in soup.find_all(True):
        for attr_val in tag.attrs.values():
            if isinstance(attr_val, str) and '@' in attr_val:
                for match in EMAIL_RE.findall(attr_val):
                    if is_valid_email(match):
                        emails.add(match.lower())

    # 4) Script tags
    for script in soup.find_all('script'):
        if script.string:
            for match in EMAIL_RE.findall(script.string):
                if is_valid_email(match):
                    emails.add(match.lower())

    # 5) JSON-LD
    for script in soup.find_all('script', type='application/ld+json'):
        if script.string:
            for match in EMAIL_RE.findall(script.string):
                if is_valid_email(match):
                    emails.add(match.lower())

    # 6) Meta tags
    for meta in soup.find_all('meta'):
        content = meta.get('content', '')
        if '@' in content:
            for match in EMAIL_RE.findall(content):
                if is_valid_email(match):
                    emails.add(match.lower())

    return emails


# --------------- Fetching ---------------

async def fetch_page(session: aiohttp.ClientSession, url: str) -> str:
    try:
        async with session.get(
            url,
            timeout=aiohttp.ClientTimeout(total=TIMEOUT),
            ssl=False,
            allow_redirects=True,
        ) as resp:
            if resp.status == 200:
                ct = resp.headers.get('Content-Type', '')
                if 'text' in ct or 'html' in ct or not ct:
                    return await resp.text(errors='replace')
    except Exception:
        pass
    return ''


def get_internal_links(html_text: str, base: str, netloc: str) -> list:
    links = set()
    soup = BeautifulSoup(html_text, 'html.parser')
    for a_tag in soup.find_all('a', href=True):
        href = a_tag['href']
        if href.startswith(('#', 'javascript:', 'tel:', 'mailto:')):
            continue
        full_url = urljoin(base, href)
        parsed = urlparse(full_url)
        if parsed.netloc == netloc:
            path_lower = parsed.path.lower()
            if any(path_lower.endswith(ext) for ext in
                   ('.pdf', '.doc', '.docx', '.zip', '.png', '.jpg',
                    '.jpeg', '.gif', '.svg', '.mp4', '.mp3')):
                continue
            clean_url = f"{parsed.scheme}://{parsed.netloc}{parsed.path}"
            links.add(clean_url)
    return list(links)


async def scrape_deep(session: aiohttp.ClientSession, url: str) -> set:
    """Deep scrape one site with a total timeout."""
    emails = set()

    if not url or url == 'None':
        return emails
    if not url.startswith('http'):
        url = 'https://' + url

    parsed = urlparse(url)
    base = f"{parsed.scheme}://{parsed.netloc}"
    netloc = parsed.netloc

    visited = set()
    to_visit = [url]
    for path in CONTACT_PATHS:
        to_visit.append(base + path)

    pages_checked = 0

    try:
        async with asyncio.timeout(SITE_TIMEOUT):
            while to_visit and pages_checked < MAX_PAGES_PER_SITE:
                current_url = to_visit.pop(0)
                if current_url in visited:
                    continue
                visited.add(current_url)

                html_text = await fetch_page(session, current_url)
                if not html_text:
                    continue

                pages_checked += 1
                emails.update(extract_emails(html_text))

                # Discover more links only if no emails yet
                if not emails and pages_checked <= 3:
                    new_links = get_internal_links(html_text, base, netloc)
                    for link in new_links:
                        if link not in visited and link not in to_visit:
                            to_visit.append(link)

                # Found emails — stop early
                if emails and pages_checked >= 2:
                    break
    except (asyncio.TimeoutError, TimeoutError):
        pass

    return emails


# --------------- Main ---------------

async def main():
    print(f"Loading {INPUT_FILE}...")
    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb.active
    total_rows = ws.max_row - 1

    rows_to_process = []
    already_have = 0
    for row_idx in range(2, ws.max_row + 1):
        email = ws.cell(row_idx, 5).value
        if email and str(email).strip():
            already_have += 1
            continue
        url = ws.cell(row_idx, 3).value or ws.cell(row_idx, 2).value or ''
        rows_to_process.append((row_idx, str(url)))

    total = len(rows_to_process)
    print(f"Total companies: {total_rows}")
    print(f"Already have emails: {already_have}")
    print(f"Need to deep-scrape: {total}\n")

    found_count = 0
    processed = 0

    connector = aiohttp.TCPConnector(limit=CONCURRENCY, ssl=False)
    headers = {
        'User-Agent': (
            'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
            'AppleWebKit/537.36 (KHTML, like Gecko) '
            'Chrome/120.0.0.0 Safari/537.36'
        ),
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
        'Accept-Language': 'ru-RU,ru;q=0.9,en;q=0.8',
    }

    async with aiohttp.ClientSession(connector=connector, headers=headers) as session:
        sem = asyncio.Semaphore(CONCURRENCY)

        async def process(row_idx: int, url: str):
            nonlocal found_count, processed
            async with sem:
                emails = await scrape_deep(session, url)
                if emails:
                    ws.cell(row_idx, 5).value = ', '.join(sorted(emails))
                    found_count += 1
                processed += 1

                if processed % 10 == 0 or processed == total:
                    pct = processed * 100 // total
                    bar = '\u2588' * (pct // 5) + '\u2591' * (20 - pct // 5)
                    print(f"\r  [{bar}] {processed}/{total}  "
                          f"new emails: {found_count}", end='', flush=True)

        tasks = [process(r, u) for r, u in rows_to_process]
        await asyncio.gather(*tasks)

    print(f"\n\n{'='*50}")
    print(f"  Deep scrape complete!")
    print(f"  Previously had emails: {already_have}")
    print(f"  New emails found:      {found_count}")
    print(f"  Total with emails:     {already_have + found_count} / {total_rows}")
    print(f"  Still missing:         {total - found_count}")
    print(f"{'='*50}")

    wb.save(OUTPUT_FILE)
    print(f"\nResults saved to: {OUTPUT_FILE}")


if __name__ == '__main__':
    start = time.time()
    asyncio.run(main())
    elapsed = time.time() - start
    m, s = divmod(int(elapsed), 60)
    print(f"Total time: {m}m {s}s")
