#!/usr/bin/env python3
"""
Unified scraper + auditor for "1000 + 86 клиник.xlsx".

For both sheets ("86" and "1000"):
  - Scrapes EMAILS for empty email cells only
  - Scrapes PHONES for empty phone cells only
  - Runs AUDIT (accessibility, theme, activity) for empty audit cells only
  - Never overwrites already filled data

Usage:
    pip install openpyxl aiohttp beautifulsoup4
    python scrape_all.py

Input:  1000 + 86 клиник.xlsx
Output: 1000 + 86 клиник_result.xlsx
"""

import asyncio
import html as html_module
import re
import time
import warnings
from datetime import datetime, timedelta
from urllib.parse import urljoin, urlparse

import aiohttp
import openpyxl
from openpyxl.styles import Font, PatternFill
from bs4 import BeautifulSoup, XMLParsedAsHTMLWarning

warnings.filterwarnings("ignore", category=XMLParsedAsHTMLWarning)
warnings.filterwarnings("ignore", category=DeprecationWarning)

# ===================== Configuration =====================

CONCURRENCY = 25
TIMEOUT = 12
SITE_TIMEOUT = 50
MAX_PAGES_PER_SITE = 12
MONTHS_BACK = 3

INPUT_FILE = '1000 + 86 клиник.xlsx'
OUTPUT_FILE = '1000 + 86 клиник_result.xlsx'

# Sheet column mappings: (url_col, email_col, phone_col, audit_start_col)
# audit columns are: status, theme, confidence, theme_detail, activity, activity_detail
SHEET_CONFIG = {
    '86': {
        'url_col': 3,        # Сайт
        'email_col': 4,      # Email
        'phone_col': None,   # No phone column
        'audit_start': 5,    # Col 5: Сайт работает?
    },
    '1000': {
        'url_col': 3,        # URL
        'email_col': 5,      # Emails
        'phone_col': 6,      # Phones
        'audit_start': 7,    # Col 7: Сайт работает?
    },
}

# ===================== Email extraction =====================

EMAIL_RE = re.compile(
    r'[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}',
    re.IGNORECASE,
)

JUNK_EMAIL_DOMAINS = {
    'example.com', 'email.com', 'domain.com', 'test.com',
    'sentry.io', 'wixpress.com', 'yoursite.com', 'yourdomain.com',
    'site.com', 'w3.org', 'schema.org', 'gravatar.com', 'wordpress.org',
    'googleusercontent.com', 'tinymce.com', 'jsdelivr.net',
    'cloudflare.com', 'googleapis.com', 'gstatic.com',
    'wordpress.com', 'wp.com', 'jquery.com', 'co.com',
    'thismywebsite.com', 'domain.ru',
}

JUNK_EMAIL_PREFIXES = (
    'noreply', 'no-reply', 'support@wix', 'support@tilda',
    'webpack', 'grunt', 'gulp', 'postmaster', 'mailer-daemon',
    'admin@wordpress', 'info@starter', 'mycompany@', 'sample@',
    'privacy@', 'coolwanglu@',
)

JUNK_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.gif', '.svg', '.webp', '.css', '.js', '.woff'}

CONTACT_PATHS = [
    '/contacts', '/contact', '/kontakty', '/about',
    '/about-us', '/o-nas', '/o-kompanii', '/kontakt',
    '/contact-us', '/svyaz', '/feedback', '/team',
    '/specialists', '/specialisty', '/komanda',
    '/o-centre', '/o-klinike', '/o-tsentre',
    '/rekvizity', '/politika-konfidencialnosti',
]


def is_valid_email(email):
    email = email.lower().strip()
    if len(email) > 100 or len(email) < 5 or '@' not in email:
        return False
    local, domain = email.rsplit('@', 1)
    if domain in JUNK_EMAIL_DOMAINS:
        return False
    if email.startswith(JUNK_EMAIL_PREFIXES):
        return False
    for ext in JUNK_EXTENSIONS:
        if email.endswith(ext) or ext in domain:
            return False
    if re.search(r'\d{5,}', email) or '..' in email:
        return False
    if not re.match(r'^[a-z0-9._%+\-]+@[a-z0-9.\-]+\.[a-z]{2,}$', email):
        return False
    return True


# ===================== Phone extraction =====================

PHONE_RE = re.compile(
    r'(?:\+7|8)[\s\-\(]*(?:\d[\s\-\)]*){10}'
)

PHONE_CLEAN_RE = re.compile(r'[^\d+]')


def extract_phones(html):
    """Extract Russian phone numbers from HTML."""
    phones = set()
    soup = BeautifulSoup(html, 'html.parser')

    # 1) tel: links
    for a_tag in soup.find_all('a', href=True):
        href = a_tag['href']
        if 'tel:' in href:
            raw = href.split('tel:')[1].split('?')[0].strip()
            cleaned = PHONE_CLEAN_RE.sub('', raw)
            if len(cleaned) >= 11:
                phones.add(format_phone(cleaned))

    # 2) Regex in text
    text = soup.get_text(' ', strip=True)
    for match in PHONE_RE.finditer(text):
        raw = match.group(0)
        cleaned = PHONE_CLEAN_RE.sub('', raw)
        if len(cleaned) >= 11:
            phones.add(format_phone(cleaned))

    # 3) data-attributes and meta
    for tag in soup.find_all(True):
        for attr_val in tag.attrs.values():
            if isinstance(attr_val, str) and ('tel:' in attr_val or re.search(r'\+7|^8\d{10}', attr_val)):
                cleaned = PHONE_CLEAN_RE.sub('', attr_val)
                if len(cleaned) >= 11:
                    phones.add(format_phone(cleaned))

    return phones


def format_phone(digits):
    """Normalize phone to +7XXXXXXXXXX format."""
    digits = digits.lstrip('+')
    if digits.startswith('8') and len(digits) == 11:
        digits = '7' + digits[1:]
    if digits.startswith('7') and len(digits) == 11:
        return f'+7{digits[1:]}'
    return f'+{digits}'


# ===================== Email extraction (deep) =====================

def extract_emails(html):
    emails = set()
    decoded = html_module.unescape(html)
    soup = BeautifulSoup(decoded, 'html.parser')

    for a_tag in soup.find_all('a', href=True):
        href = a_tag['href']
        if 'mailto:' in href:
            raw = href.split('mailto:')[1].split('?')[0].split('#')[0].strip()
            if is_valid_email(raw):
                emails.add(raw.lower())

    for match in EMAIL_RE.findall(decoded):
        if is_valid_email(match):
            emails.add(match.lower())

    for tag in soup.find_all(True):
        for attr_val in tag.attrs.values():
            if isinstance(attr_val, str) and '@' in attr_val:
                for match in EMAIL_RE.findall(attr_val):
                    if is_valid_email(match):
                        emails.add(match.lower())

    for script in soup.find_all('script'):
        if script.string:
            for match in EMAIL_RE.findall(script.string):
                if is_valid_email(match):
                    emails.add(match.lower())

    for script in soup.find_all('script', type='application/ld+json'):
        if script.string:
            for match in EMAIL_RE.findall(script.string):
                if is_valid_email(match):
                    emails.add(match.lower())

    for meta in soup.find_all('meta'):
        content = meta.get('content', '')
        if '@' in content:
            for match in EMAIL_RE.findall(content):
                if is_valid_email(match):
                    emails.add(match.lower())

    return emails


# ===================== Theme checking =====================

PRIMARY_KEYWORDS = [
    'детский психолог', 'детская психолог', 'детских психолог',
    'нейропсихолог', 'нейро-психолог',
    'детский центр', 'детский развива',
    'логопед', 'дефектолог',
    'детский нейро', 'нейрокоррекц',
    'психолог для детей', 'психолог для ребенк', 'психолог для ребёнк',
    'психологическ центр', 'психологический центр',
    'коррекционн', 'коррекция',
    'развити ребенк', 'развити ребёнк', 'развитие детей',
    'aba-терап', 'аба-терап', 'аба терап',
    'сенсорная интеграц', 'канистерап',
    'детский невролог', 'детский психиатр',
    'песочная терап', 'арт-терап', 'игровая терап',
    'задержк развит', 'зпр', 'зрр', 'сдвг', 'аутизм', 'рас ',
    'дислекси', 'дисграф', 'алали', 'заикани',
]

SECONDARY_KEYWORDS = [
    'психолог', 'психотерап', 'консультац', 'тревожност', 'эмоциональн',
]

ANTI_KEYWORDS = [
    'автосервис', 'автомобил', 'ремонт квартир', 'натяжн потолк',
    'юридическ', 'бухгалтер', 'строительств', 'недвижимост',
    'стоматолог', 'кулинар', 'доставка еды',
]


def check_theme(html):
    if not html:
        return None, 0, 'Нет данных'
    text = BeautifulSoup(html, 'html.parser').get_text(' ', strip=True).lower()
    anti_count = sum(1 for kw in ANTI_KEYWORDS if kw in text)
    if anti_count >= 2:
        return False, 0, 'Другая тематика'
    primary_found = [kw for kw in PRIMARY_KEYWORDS if kw in text]
    secondary_found = [kw for kw in SECONDARY_KEYWORDS if kw in text]
    if len(primary_found) >= 3:
        return True, 100, f'Точное совпадение ({len(primary_found)} ключевых слов)'
    elif len(primary_found) >= 1:
        conf = min(50 + len(primary_found) * 20 + len(secondary_found) * 10, 100)
        return True, conf, f'Совпадение ({len(primary_found)} осн. + {len(secondary_found)} доп.)'
    elif len(secondary_found) >= 2:
        return True, 40, f'Возможное совпадение ({len(secondary_found)} доп. слов)'
    elif len(secondary_found) == 1:
        return None, 20, 'Слабое совпадение (1 общее слово)'
    return False, 0, 'Не соответствует тематике'


# ===================== Activity checking =====================

RUSSIAN_MONTHS = {
    'январ': 1, 'феврал': 2, 'март': 3, 'апрел': 4,
    'мая': 5, 'мае': 5, 'май': 5, 'июн': 6, 'июл': 7,
    'август': 8, 'сентябр': 9, 'октябр': 10, 'ноябр': 11, 'декабр': 12,
}

DATE_PATTERNS = [
    re.compile(r'(\d{1,2})\s+(' + '|'.join(RUSSIAN_MONTHS.keys()) + r')[а-яё]*\s+(\d{4})', re.IGNORECASE),
    re.compile(r'(\d{1,2})[./](\d{1,2})[./](\d{4})'),
    re.compile(r'(\d{4})-(\d{2})-(\d{2})'),
]

COPYRIGHT_RE = re.compile(r'©\s*(?:\d{4}\s*[-–—]\s*)?(\d{4})', re.IGNORECASE)


def parse_date(match, idx):
    try:
        if idx == 0:
            day, month_text, year = int(match.group(1)), match.group(2).lower(), int(match.group(3))
            month = next((v for k, v in RUSSIAN_MONTHS.items() if month_text.startswith(k)), None)
            if month and 2000 <= year <= 2030:
                return datetime(year, month, day)
        elif idx == 1:
            d, m, y = int(match.group(1)), int(match.group(2)), int(match.group(3))
            if 2000 <= y <= 2030 and 1 <= m <= 12:
                return datetime(y, m, d)
        elif idx == 2:
            y, m, d = int(match.group(1)), int(match.group(2)), int(match.group(3))
            if 2000 <= y <= 2030 and 1 <= m <= 12:
                return datetime(y, m, d)
    except (ValueError, IndexError):
        pass
    return None


def check_activity(html, headers):
    now = datetime.now()
    cutoff = now - timedelta(days=MONTHS_BACK * 30)
    results = []

    lm = headers.get('Last-Modified', '')
    if lm:
        try:
            from email.utils import parsedate_to_datetime
            lm_date = parsedate_to_datetime(lm).replace(tzinfo=None)
            if lm_date >= cutoff:
                results.append(f'Last-Modified: {lm_date.strftime("%d.%m.%Y")}')
        except Exception:
            pass

    if not html:
        return (True, '; '.join(results)) if results else (None, 'Нет данных')

    text = BeautifulSoup(html, 'html.parser').get_text(' ', strip=True)

    dates = []
    for i, pat in enumerate(DATE_PATTERNS):
        for m in pat.finditer(text):
            dt = parse_date(m, i)
            if dt:
                dates.append(dt)

    recent = [d for d in dates if cutoff <= d <= now + timedelta(days=30)]
    if recent:
        latest = max(recent)
        results.append(f'Дата на сайте: {latest.strftime("%d.%m.%Y")} ({len(recent)} свежих)')

    cr = COPYRIGHT_RE.search(text)
    if cr:
        yr = int(cr.group(1))
        if yr >= now.year:
            results.append(f'Copyright {yr}')
        elif yr == now.year - 1:
            results.append(f'Copyright {yr} (прошлый год)')

    if results:
        return True, '; '.join(results)
    if dates:
        return False, f'Последняя дата: {max(dates).strftime("%d.%m.%Y")}'
    if cr:
        yr = int(cr.group(1))
        if yr < now.year - 1:
            return False, f'Copyright {yr} (устаревший)'
    return None, 'Нет признаков активности'


# ===================== Fetching =====================

async def fetch_page(session, url):
    try:
        async with session.get(
            url, timeout=aiohttp.ClientTimeout(total=TIMEOUT),
            ssl=False, allow_redirects=True,
        ) as resp:
            if resp.status == 200:
                ct = resp.headers.get('Content-Type', '')
                if 'text' in ct or 'html' in ct or not ct:
                    return resp.status, await resp.text(errors='replace'), dict(resp.headers)
            return resp.status, '', dict(resp.headers)
    except aiohttp.ClientConnectorError:
        return 0, '', {}
    except asyncio.TimeoutError:
        return -1, '', {}
    except Exception:
        return -2, '', {}


async def fetch_contact_pages(session, base_url, main_html, netloc):
    """Fetch contact-like pages and return combined HTML."""
    extra_html = []
    checked = set()

    # Links found on main page
    if main_html:
        soup = BeautifulSoup(main_html, 'html.parser')
        keywords = ['contact', 'kontakt', 'контакт', 'связ', 'о нас', 'about', 'обратн', 'напис']
        for a in soup.find_all('a', href=True):
            href_lower = a['href'].lower()
            text_lower = (a.get_text() or '').lower()
            if any(kw in href_lower or kw in text_lower for kw in keywords):
                full = urljoin(base_url, a['href'])
                if urlparse(full).netloc == netloc and full not in checked:
                    checked.add(full)
                    _, h, _ = await fetch_page(session, full)
                    if h:
                        extra_html.append(h)
                    if len(checked) >= 5:
                        break

    # Common paths
    for path in CONTACT_PATHS[:10]:
        curl = base_url + path
        if curl not in checked:
            checked.add(curl)
            _, h, _ = await fetch_page(session, curl)
            if h:
                extra_html.append(h)

    return extra_html


def status_text(code):
    if code == 200: return 'OK'
    if code == 0: return 'Не открывается (ошибка соединения)'
    if code == -1: return 'Не открывается (таймаут)'
    if code == -2: return 'Не открывается (ошибка)'
    if code == 403: return 'Заблокирован (403)'
    if code == 404: return 'Не найден (404)'
    if 500 <= code < 600: return f'Ошибка сервера ({code})'
    if 300 <= code < 400: return f'Редирект ({code})'
    return f'Ошибка ({code})'


# ===================== Process one row =====================

async def process_row(session, url, need_email, need_phone, need_audit):
    """Process a single site. Returns (emails, phones, audit_data)."""
    emails = set()
    phones = set()
    audit = None

    if not url or url == 'None' or url == 'EMPTY':
        return None, None, None

    if not url.startswith('http'):
        url = 'https://' + url

    parsed = urlparse(url)
    base = f"{parsed.scheme}://{parsed.netloc}"
    netloc = parsed.netloc

    # Fetch main page
    try:
        async with asyncio.timeout(SITE_TIMEOUT):
            status_code, main_html, headers = await fetch_page(session, url)

            all_html = [main_html] if main_html else []

            # Fetch contact pages if we need email or phone
            if main_html and (need_email or need_phone):
                extra = await fetch_contact_pages(session, base, main_html, netloc)
                all_html.extend(extra)

            # Extract emails
            if need_email:
                for h in all_html:
                    emails.update(extract_emails(h))

            # Extract phones
            if need_phone:
                for h in all_html:
                    phones.update(extract_phones(h))

            # Audit
            if need_audit:
                st = status_text(status_code)
                theme, theme_conf, theme_detail = check_theme(main_html)
                active, active_detail = check_activity(main_html, headers)
                audit = (st, status_code, theme, theme_conf, theme_detail, active, active_detail)

    except (asyncio.TimeoutError, TimeoutError):
        if need_audit:
            audit = ('Не открывается (таймаут)', -1, None, 0, 'Нет данных', None, 'Нет данных')

    return emails or None, phones or None, audit


# ===================== Main =====================

GREEN = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
RED = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
YELLOW = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')


async def process_sheet(session, ws, config, sheet_name):
    """Process one sheet."""
    total = ws.max_row - 1
    if total <= 0:
        print(f"  Sheet '{sheet_name}': empty, skipping")
        return

    url_col = config['url_col']
    email_col = config['email_col']
    phone_col = config['phone_col']
    audit_start = config['audit_start']

    # Count what needs processing
    rows_to_process = []
    skip_count = 0
    for row_idx in range(2, ws.max_row + 1):
        url = ws.cell(row_idx, url_col).value
        if not url or not str(url).strip():
            skip_count += 1
            continue

        url = str(url).strip()

        need_email = email_col and (not ws.cell(row_idx, email_col).value or
                                     not str(ws.cell(row_idx, email_col).value).strip())
        need_phone = phone_col and (not ws.cell(row_idx, phone_col).value or
                                     not str(ws.cell(row_idx, phone_col).value).strip())
        # Audit: check if "Сайт работает?" column is empty
        need_audit = not ws.cell(row_idx, audit_start).value or \
                     not str(ws.cell(row_idx, audit_start).value).strip()

        if not need_email and not need_phone and not need_audit:
            skip_count += 1
            continue

        rows_to_process.append((row_idx, url, need_email, need_phone, need_audit))

    to_process = len(rows_to_process)
    print(f"  Sheet '{sheet_name}': {total} rows, {to_process} to process, {skip_count} already filled")

    if to_process == 0:
        return

    processed = 0
    email_found = 0
    phone_found = 0

    sem = asyncio.Semaphore(CONCURRENCY)

    async def do_row(row_idx, url, need_email, need_phone, need_audit):
        nonlocal processed, email_found, phone_found
        async with sem:
            emails, phones, audit = await process_row(
                session, url, need_email, need_phone, need_audit
            )

            if emails and need_email:
                ws.cell(row_idx, email_col).value = ', '.join(sorted(emails))
                email_found += 1

            if phones and need_phone:
                ws.cell(row_idx, phone_col).value = ', '.join(sorted(phones))
                phone_found += 1

            if audit and need_audit:
                st, st_code, theme, theme_conf, theme_det, active, active_det = audit
                ac = audit_start

                ws.cell(row_idx, ac).value = st
                if st_code == 200:
                    ws.cell(row_idx, ac).fill = GREEN
                elif st_code > 0:
                    ws.cell(row_idx, ac).fill = YELLOW
                else:
                    ws.cell(row_idx, ac).fill = RED

                if theme is True:
                    ws.cell(row_idx, ac+1).value = 'Да'
                    ws.cell(row_idx, ac+1).fill = GREEN
                elif theme is False:
                    ws.cell(row_idx, ac+1).value = 'Нет'
                    ws.cell(row_idx, ac+1).fill = RED
                else:
                    ws.cell(row_idx, ac+1).value = 'Неясно'
                    ws.cell(row_idx, ac+1).fill = YELLOW

                ws.cell(row_idx, ac+2).value = theme_conf
                ws.cell(row_idx, ac+3).value = theme_det

                if active is True:
                    ws.cell(row_idx, ac+4).value = 'Активен'
                    ws.cell(row_idx, ac+4).fill = GREEN
                elif active is False:
                    ws.cell(row_idx, ac+4).value = 'Неактивен'
                    ws.cell(row_idx, ac+4).fill = RED
                else:
                    ws.cell(row_idx, ac+4).value = 'Неизвестно'
                    ws.cell(row_idx, ac+4).fill = YELLOW

                ws.cell(row_idx, ac+5).value = active_det

            processed += 1
            if processed % 15 == 0 or processed == to_process:
                pct = processed * 100 // to_process
                bar = '\u2588' * (pct // 5) + '\u2591' * (20 - pct // 5)
                print(f"\r  [{bar}] {processed}/{to_process}  "
                      f"emails: +{email_found}  phones: +{phone_found}",
                      end='', flush=True)

    tasks = [do_row(r, u, ne, np_, na) for r, u, ne, np_, na in rows_to_process]
    await asyncio.gather(*tasks)

    print(f"\n  => Done! Emails: +{email_found}, Phones: +{phone_found}\n")


async def main():
    print(f"Loading {INPUT_FILE}...\n")
    wb = openpyxl.load_workbook(INPUT_FILE)

    connector = aiohttp.TCPConnector(limit=CONCURRENCY, ssl=False)
    headers = {
        'User-Agent': (
            'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
            'AppleWebKit/537.36 (KHTML, like Gecko) '
            'Chrome/120.0.0.0 Safari/537.36'
        ),
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
        'Accept-Language': 'ru-RU,ru;q=0.9,en;q=0.8',
    }

    async with aiohttp.ClientSession(connector=connector, headers=headers) as session:
        for sheet_name in wb.sheetnames:
            if sheet_name in SHEET_CONFIG:
                print(f"Processing sheet: {sheet_name}")
                ws = wb[sheet_name]
                await process_sheet(session, ws, SHEET_CONFIG[sheet_name], sheet_name)
            else:
                print(f"Skipping unknown sheet: {sheet_name}")

    wb.save(OUTPUT_FILE)
    print(f"Results saved to: {OUTPUT_FILE}")


if __name__ == '__main__':
    start = time.time()
    asyncio.run(main())
    elapsed = time.time() - start
    m, s = divmod(int(elapsed), 60)
    print(f"Total time: {m}m {s}s")
